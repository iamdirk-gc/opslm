"""Spreadsheet-based labeling workflow for evaluation cases.

Reviewers work in Excel or Google Sheets, not in code:

1. ``plan_batches`` / ``export_review_sheet``: candidate pairs -> a sheet per reviewer.
2. Each reviewer records the two guideline judgments per row (the label is derived, never typed),
   the reasons, and the evidence spans that support the label.
3. ``import_review_sheet``: validate the sheet and turn it into ``Review`` records. Problems come
   back as plain-language messages with sheet row numbers.
4. ``compare_reviews``: agreement (Cohen's kappa) on the shared overlap and a disagreement sheet
   for the adjudication meeting, where the deciding rule is written down.
5. ``finalize``: reviews plus resolutions -> ``EvalCase`` records (input + gold label), validated
   for leakage.

The sheet is only a transport. Reviewers' edits to the read-only columns are ignored on import
because rows are matched on ``case_id`` and facts come from the candidate file.
"""

from __future__ import annotations

import csv
import hashlib
import random
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

from pydantic import BaseModel, ConfigDict

from opslm.data import validate_cases
from opslm.labels import GUIDELINE_VERSION, derive_label, values_agree
from opslm.schema import (
    BusinessFact,
    EvalCase,
    GoldLabel,
    ReasonCode,
    ReconciliationInput,
    Relationship,
    Source,
    Split,
    reasons_problem,
)

VERDICTS = ("yes", "no", "cannot_tell")
RESOLVER = "adjudicated"


class _Model(BaseModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


class CandidatePair(_Model):
    """An unlabeled pair awaiting review, with the metadata the final case will carry."""

    input: ReconciliationInput
    source: Source
    split: Split
    tags: tuple[str, ...] = ()

    @property
    def case_id(self) -> str:
        return self.input.case_id


class Review(_Model):
    """One reviewer's judgments for one pair. The label is derived from these, not stored."""

    case_id: str
    reviewer: str
    concept_same: bool | None
    scope_same: bool | None
    reasons: tuple[ReasonCode, ...]  # primary first
    evidence_refs: tuple[str, ...]
    rationale: str
    flag: str = ""


def read_candidates(path: str | Path) -> list[CandidatePair]:
    with open(path, encoding="utf-8") as f:
        return [CandidatePair.model_validate_json(line) for line in f if line.strip()]


def write_candidates(path: str | Path, candidates: Iterable[CandidatePair]) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.writelines(c.model_dump_json() + "\n" for c in sorted(candidates, key=lambda c: c.case_id))


def read_reviews(path: str | Path) -> list[Review]:
    with open(path, encoding="utf-8") as f:
        return [Review.model_validate_json(line) for line in f if line.strip()]


def write_reviews(path: str | Path, reviews: Iterable[Review]) -> None:
    with open(path, "w", encoding="utf-8") as f:
        f.writelines(r.model_dump_json() + "\n" for r in sorted(reviews, key=lambda r: r.case_id))


def _values_agree(c: CandidatePair) -> bool | None:
    return values_agree(c.input.fact_a.quantity, c.input.fact_b.quantity)


def derived_label(candidate: CandidatePair, review: Review) -> Relationship:
    return derive_label(review.concept_same, review.scope_same, _values_agree(candidate))


def _family(c: CandidatePair) -> str:
    return ", ".join(sorted(c.input.document_families()))


# --- sheet layout ----------------------------------------------------------------

_FACT_COLS = ("metric", "metric_id", "value", "entity", "period", "definition", "basis",
              "source", "evidence")
READONLY = ["case_id", "document_family", "values_agree"] + [
    f"{side}_{c}" for side in ("A", "B") for c in _FACT_COLS
]
REVIEWER_COLS = ["concept_same", "scope_same", "primary_reason", "other_reasons",
                 "evidence_refs", "rationale", "flag"]
REVIEW_COLUMNS = READONLY + REVIEWER_COLS
REASONS = [r.value for r in ReasonCode]
DROPDOWNS = {
    "concept_same": list(VERDICTS),
    "scope_same": list(VERDICTS),
    "primary_reason": REASONS,
}
RESOLUTION_COLS = ["resolved_concept_same", "resolved_scope_same", "resolved_primary_reason",
                   "resolved_other_reasons", "resolved_evidence_refs", "resolved_rationale",
                   "rule_decided"]
RESOLUTION_DROPDOWNS = {
    "resolved_concept_same": list(VERDICTS),
    "resolved_scope_same": list(VERDICTS),
    "resolved_primary_reason": REASONS,
}
_REVIEW_PREFIXES = ("r1", "r2")
DISAGREEMENT_COLUMNS = (
    ["case_id", "values_agree", "why"]
    + [f"{side}_{c}" for side in ("A", "B") for c in _FACT_COLS]
    + [f"{p}_{c}" for p in _REVIEW_PREFIXES for c in ["reviewer"] + REVIEWER_COLS]
    + RESOLUTION_COLS
)


def _period_text(fact: BusinessFact) -> str:
    p = fact.period
    if p is None:
        return ""
    if p.type == "instant":
        return f"{p.label} [as of {p.as_of or '?'}]"
    details = [f"FY{p.fiscal_year}" if p.fiscal_year else "",
               f"Q{p.fiscal_quarter}" if p.fiscal_quarter else "", "LTM" if p.is_ltm else ""]
    if p.start or p.end:
        details.append(f"{p.start or '?'} to {p.end or '?'}")
    shown = " ".join(d for d in details if d)
    return f"{p.label} [{shown}]" if shown else p.label


def _value_text(fact: BusinessFact) -> str:
    q = fact.quantity
    unit = q.currency if q.unit == "currency" else q.unit
    return f"{q.raw_text} (= {q.value:,f} {unit})"


def _entity_text(fact: BusinessFact) -> str:
    e = fact.entity
    if e is None:
        return ""
    return e.name + (f" [{e.scope}" + (f": {e.segment}" if e.segment else "") + "]"
                     if e.scope != "consolidated" or e.segment else "")


def _fact_cells(side: str, fact: BusinessFact, doc_types: dict[str, str]) -> dict[str, str]:
    first = fact.evidence[0]
    loc = first.cell_ref or (f"p.{first.page}" if first.page else "")
    return {
        f"{side}_metric": fact.metric_label,
        f"{side}_metric_id": fact.metric_id,
        f"{side}_value": _value_text(fact),
        f"{side}_entity": _entity_text(fact),
        f"{side}_period": _period_text(fact),
        f"{side}_definition": fact.stated_definition or "",
        f"{side}_basis": f"{fact.accounting_basis}, {fact.reporting_status}",
        f"{side}_source": " ".join(x for x in (doc_types.get(first.doc_id), loc) if x),
        f"{side}_evidence": "\n".join(f"[{s.span_id}] {s.text}" for s in fact.evidence),
    }


def _candidate_cells(c: CandidatePair) -> dict[str, str]:
    doc_types = {d.doc_id: d.doc_type for d in c.input.documents}
    return {**_fact_cells("A", c.input.fact_a, doc_types),
            **_fact_cells("B", c.input.fact_b, doc_types)}


def _agree_text(c: CandidatePair) -> str:
    agree = _values_agree(c)
    return "not comparable" if agree is None else "yes" if agree else "no"


def review_rows(candidates: Iterable[CandidatePair]) -> list[dict[str, str]]:
    rows = []
    for c in candidates:
        row = {"case_id": c.case_id, "document_family": _family(c), "values_agree": _agree_text(c)}
        row.update(_candidate_cells(c))
        row.update(dict.fromkeys(REVIEWER_COLS, ""))
        rows.append(row)
    return rows


# --- file I/O (csv and xlsx) -------------------------------------------------------

_FORMULA_START = re.compile(r"^(?:[=+@\t\r]|-(?![\d.]))")


def _safe(text: str) -> str:
    """Stop spreadsheet apps from running document text as a formula."""
    return "'" + text if _FORMULA_START.match(text) else text


def write_sheet(
    path: str | Path,
    rows: list[dict[str, str]],
    columns: list[str],
    *,
    dropdowns: dict[str, list[str]] | None = None,
    editable: Iterable[str] = (),
) -> None:
    """Write rows as .csv or .xlsx (xlsx adds dropdowns, frozen header and shading)."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(columns)
            for row in rows:
                w.writerow([_safe(row.get(c, "")) for c in columns])
        return
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "review"
    ws.append(columns)
    for row in rows:
        ws.append([_safe(row.get(c, "")) for c in columns])
    for row in ws.iter_rows(min_row=2):  # force text so nothing is parsed as a formula or number
        for cell in row:
            cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    edit = set(editable)
    grey, yellow = PatternFill("solid", fgColor="EEEEEE"), PatternFill("solid", fgColor="FFF8DC")
    for i, name in enumerate(columns, start=1):
        letter = get_column_letter(i)
        head = ws.cell(row=1, column=1 + i - 1)
        head.font = Font(bold=True)
        head.fill = yellow if name in edit else grey
        ws.column_dimensions[letter].width = 38 if name.endswith(("excerpt", "rationale", "definition")) else 18
        if name in (dropdowns or {}):
            dv = DataValidation(type="list", formula1='"' + ",".join(dropdowns[name]) + '"',
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{max(len(rows) + 1, 2)}")
    ws.freeze_panes = "B2"
    wb.save(path)


def read_sheet(path: str | Path) -> list[dict[str, str]]:
    """Read a .csv or .xlsx sheet into dicts keyed by header. Missing cells become ''."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [{k: (v or "").strip() for k, v in row.items() if k} for row in csv.DictReader(f)]
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it, ())]
    return [
        {h: ("" if v is None else str(v).strip()) for h, v in zip(header, row, strict=False) if h}
        for row in it
    ]


def export_review_sheet(path: str | Path, candidates: Iterable[CandidatePair]) -> None:
    write_sheet(path, review_rows(candidates), REVIEW_COLUMNS, dropdowns=DROPDOWNS,
                editable=REVIEWER_COLS)


# --- import and validation ----------------------------------------------------------


class ImportResult(_Model):
    reviews: list[Review]
    problems: list[str]
    unlabeled: list[str]


_VERDICT_ALIASES = {"yes": True, "y": True, "no": False, "n": False,
                    "cannot_tell": None, "cannot tell": None, "cannot-tell": None, "unknown": None}


def _verdict(text: str) -> tuple[bool | None, bool]:
    """Return (value, ok). Blank is not ok."""
    key = text.strip().lower()
    return (_VERDICT_ALIASES[key], True) if key in _VERDICT_ALIASES else (None, False)


def _items(text: str) -> list[str]:
    return [t.strip() for t in re.split(r"[,;\s]+", text) if t.strip()]


def _parse_judgments(
    row: dict[str, str], prefix: str, candidate: CandidatePair, where: str
) -> tuple[dict[str, Any] | None, list[str]]:
    """Shared validation for review rows (prefix '') and resolution rows ('resolved_')."""
    col = lambda name: row.get(f"{prefix}{name}", "").strip()
    problems: list[str] = []
    verdicts = {}
    for name in ("concept_same", "scope_same"):
        value, ok = _verdict(col(name))
        if not ok:
            shown = col(name) or "blank"
            problems.append(f"{where}: {prefix}{name} is {shown!r}; use yes, no or cannot_tell")
        verdicts[name] = value
    if problems:
        return None, problems

    label = derive_label(verdicts["concept_same"], verdicts["scope_same"], _values_agree(candidate))
    primary = col("primary_reason").lower()
    reasons = ([primary] if primary else []) + [r.lower() for r in _items(col("other_reasons"))]
    if not primary:
        problems.append(f"{where}: your judgments give '{label.value}'; {prefix}primary_reason "
                        "is required")
    bad = [r for r in reasons if r not in REASONS]
    if bad:
        problems.append(f"{where}: unknown reason {', '.join(map(repr, bad))}; see the "
                        "primary_reason dropdown for valid codes")
    elif primary:
        misfit = reasons_problem(label.value, reasons)
        if misfit:
            problems.append(f"{where}: your judgments give '{label.value}', but {misfit}")

    refs = _items(col("evidence_refs"))
    unknown = [r for r in refs if r not in candidate.input.span_ids()]
    if unknown:
        problems.append(f"{where}: unknown evidence_refs {unknown}; this pair's spans are "
                        f"{', '.join(sorted(candidate.input.span_ids()))}")
    if label != Relationship.INSUFFICIENT_EVIDENCE and not refs and not unknown:
        problems.append(f"{where}: evidence_refs is required for '{label.value}' (list the span "
                        "ids shown in brackets in the evidence columns)")
    if not col("rationale"):
        problems.append(f"{where}: rationale is required")
    if problems:
        return None, problems
    return {"concept_same": verdicts["concept_same"], "scope_same": verdicts["scope_same"],
            "reasons": tuple(dict.fromkeys(reasons)), "evidence_refs": tuple(refs),
            "rationale": col("rationale")}, []


def import_review_sheet(
    rows: list[dict[str, str]], candidates: Iterable[CandidatePair], reviewer: str
) -> ImportResult:
    by_id = {c.case_id: c for c in candidates}
    reviews: list[Review] = []
    problems: list[str] = []
    unlabeled: list[str] = []
    seen: set[str] = set()
    for n, row in enumerate(rows, start=2):  # row 1 is the header
        case_id = row.get("case_id", "").strip()
        if case_id not in by_id:
            problems.append(f"row {n}: unknown case_id {case_id!r}")
            continue
        if case_id in seen:
            problems.append(f"row {n} ({case_id}): case appears more than once")
            continue
        seen.add(case_id)
        if not any(row.get(c, "").strip() for c in REVIEWER_COLS if c != "flag"):
            unlabeled.append(case_id)
            continue
        parsed, errs = _parse_judgments(row, "", by_id[case_id], f"row {n} ({case_id})")
        problems += errs
        if parsed:
            reviews.append(Review(case_id=case_id, reviewer=reviewer, flag=row.get("flag", ""),
                                  **parsed))
    return ImportResult(reviews=reviews, problems=problems, unlabeled=unlabeled)


# --- batches ------------------------------------------------------------------------


def _rank(case_id: str, seed: int) -> str:
    return hashlib.sha256(f"{seed}:{case_id}".encode()).hexdigest()


def plan_batches(
    candidates: Iterable[CandidatePair], reviewers: list[str], *, overlap: int, seed: int = 0
) -> dict[str, list[str]]:
    """Assign case ids to reviewers: a shared overlap, the rest split evenly.

    Deterministic for a given seed. Each reviewer's list is shuffled independently so order
    effects do not line up across reviewers.
    """
    ids = sorted((c.case_id for c in candidates), key=lambda i: _rank(i, seed))
    if overlap > len(ids):
        raise ValueError(f"overlap {overlap} exceeds the {len(ids)} candidates")
    if len(reviewers) < 2 or len(set(reviewers)) != len(reviewers):
        raise ValueError("need at least two distinct reviewers")
    shared, rest = ids[:overlap], ids[overlap:]
    plan = {r: list(shared) + rest[i :: len(reviewers)] for i, r in enumerate(reviewers)}
    for r, batch in plan.items():
        random.Random(f"{seed}:{r}").shuffle(batch)
    return plan


# --- agreement ----------------------------------------------------------------------


def cohen_kappa(x: list[str], y: list[str]) -> float | None:
    """Cohen's kappa for two raters. None when chance agreement is total (undefined)."""
    n = len(x)
    if n == 0 or n != len(y):
        return None
    po = sum(a == b for a, b in zip(x, y, strict=True)) / n
    labels = set(x) | set(y)
    pe = sum((x.count(k) / n) * (y.count(k) / n) for k in labels)
    return None if pe == 1 else (po - pe) / (1 - pe)


def _vtext(v: bool | None) -> str:
    return "cannot_tell" if v is None else "yes" if v else "no"


class Agreement(_Model):
    n: int
    label_agreement: float
    label_kappa: float | None
    concept_agreement: float
    concept_kappa: float | None
    scope_agreement: float
    scope_kappa: float | None
    disagreements: list[str]
    reason_differs: list[str]  # judgments agree but the primary reason differs (needs adjudication)
    evidence_differs: list[str]  # judgments and primary reason agree but the cited spans differ
    flagged: list[str]


def compare_reviews(
    a: Iterable[Review], b: Iterable[Review], candidates: Iterable[CandidatePair]
) -> Agreement:
    """Agreement on the cases both reviewers did. A case is a disagreement if the final label
    or either judgment differs. Two milder differences are reported separately: a different
    primary reason (``reason_differs``, which needs adjudication because results are sliced by it)
    and different cited spans with everything else equal (``evidence_differs``, which ``finalize``
    resolves by taking the union and flagging it, since reviewers often cite different numbers of
    spans for the same conclusion)."""
    by_id = {c.case_id: c for c in candidates}
    ra, rb = {r.case_id: r for r in a}, {r.case_id: r for r in b}
    shared = sorted(set(ra) & set(rb))
    lab = ([derived_label(by_id[i], ra[i]).value for i in shared],
           [derived_label(by_id[i], rb[i]).value for i in shared])
    con = ([_vtext(ra[i].concept_same) for i in shared], [_vtext(rb[i].concept_same) for i in shared])
    sco = ([_vtext(ra[i].scope_same) for i in shared], [_vtext(rb[i].scope_same) for i in shared])

    def agree(pair: tuple[list[str], list[str]]) -> float:
        return sum(p == q for p, q in zip(*pair, strict=True)) / len(shared) if shared else 0.0

    disagreements = [
        i for k, i in enumerate(shared)
        if lab[0][k] != lab[1][k] or con[0][k] != con[1][k] or sco[0][k] != sco[1][k]
    ]
    reason_differs = [i for i in shared if i not in disagreements and not _same_reason(ra[i], rb[i])]
    evidence_differs = [i for i in shared if i not in disagreements and i not in reason_differs
                        and not _same_evidence(ra[i], rb[i])]
    return Agreement(
        n=len(shared),
        label_agreement=agree(lab), label_kappa=cohen_kappa(*lab),
        concept_agreement=agree(con), concept_kappa=cohen_kappa(*con),
        scope_agreement=agree(sco), scope_kappa=cohen_kappa(*sco),
        disagreements=disagreements,
        reason_differs=reason_differs,
        evidence_differs=evidence_differs,
        flagged=[i for i in shared if ra[i].flag or rb[i].flag],
    )


def disagreement_rows(
    a: Iterable[Review], b: Iterable[Review], candidates: Iterable[CandidatePair],
    agreement: Agreement,
) -> list[dict[str, str]]:
    by_id = {c.case_id: c for c in candidates}
    ra, rb = {r.case_id: r for r in a}, {r.case_id: r for r in b}
    rows = []
    needs_row = set(agreement.disagreements) | set(agreement.reason_differs) | set(agreement.flagged)
    for case_id in sorted(needs_row):
        c = by_id[case_id]
        why = ("disagree" if case_id in agreement.disagreements
               else "primary reason differs" if case_id in agreement.reason_differs else "flagged")
        row = {"case_id": case_id, "values_agree": _agree_text(c), "why": why}
        row.update(_candidate_cells(c))
        for prefix, review in zip(_REVIEW_PREFIXES, (ra[case_id], rb[case_id]), strict=True):
            row[f"{prefix}_reviewer"] = review.reviewer
            row[f"{prefix}_concept_same"] = _vtext(review.concept_same)
            row[f"{prefix}_scope_same"] = _vtext(review.scope_same)
            row[f"{prefix}_primary_reason"] = review.reasons[0] if review.reasons else ""
            row[f"{prefix}_other_reasons"] = ", ".join(review.reasons[1:])
            row[f"{prefix}_evidence_refs"] = ", ".join(review.evidence_refs)
            row[f"{prefix}_rationale"] = review.rationale
            row[f"{prefix}_flag"] = review.flag
        row.update(dict.fromkeys(RESOLUTION_COLS, ""))
        rows.append(row)
    return rows


def export_disagreement_sheet(
    path: str | Path, a: list[Review], b: list[Review], candidates: list[CandidatePair],
    agreement: Agreement,
) -> None:
    write_sheet(path, disagreement_rows(a, b, candidates, agreement), DISAGREEMENT_COLUMNS,
                dropdowns=RESOLUTION_DROPDOWNS, editable=RESOLUTION_COLS)


def import_resolutions(
    rows: list[dict[str, str]], candidates: Iterable[CandidatePair]
) -> ImportResult:
    """Read the adjudication sheet. Every resolved row must state the rule that decided it."""
    by_id = {c.case_id: c for c in candidates}
    reviews: list[Review] = []
    problems: list[str] = []
    unlabeled: list[str] = []
    for n, row in enumerate(rows, start=2):
        case_id = row.get("case_id", "").strip()
        if case_id not in by_id:
            problems.append(f"row {n}: unknown case_id {case_id!r}")
            continue
        if not any(row.get(c, "").strip() for c in RESOLUTION_COLS):
            unlabeled.append(case_id)
            continue
        where = f"row {n} ({case_id})"
        parsed, errs = _parse_judgments(row, "resolved_", by_id[case_id], where)
        if not row.get("rule_decided", "").strip():
            errs.append(f"{where}: rule_decided is required (write down the rule that settled it)")
        problems += errs
        if parsed and not errs:
            reviews.append(Review(case_id=case_id, reviewer=RESOLVER,
                                  flag=row["rule_decided"].strip(), **parsed))
    return ImportResult(reviews=reviews, problems=problems, unlabeled=unlabeled)


# --- finalize -----------------------------------------------------------------------


class FinalizeResult(_Model):
    cases: list[EvalCase]
    problems: list[str]
    unlabeled: list[str]
    evidence_unioned: list[str] = []  # reviewers cited different spans; the union was taken


def _same_judgments(x: Review, y: Review) -> bool:
    return (x.concept_same, x.scope_same) == (y.concept_same, y.scope_same)


def _same_reason(x: Review, y: Review) -> bool:
    return x.reasons[:1] == y.reasons[:1]


def _same_evidence(x: Review, y: Review) -> bool:
    return set(x.evidence_refs) == set(y.evidence_refs)


def finalize(
    candidates: Iterable[CandidatePair],
    review_sets: list[list[Review]],
    resolutions: Iterable[Review] = (),
) -> FinalizeResult:
    """Turn reviews into evaluation cases.

    A resolution wins. Otherwise reviewers must agree on both judgments and on the primary
    reason (the first reviewer's reasons and rationale are kept; the second is recorded as
    ``reviewed_by``). If they cite different evidence spans, the gold label takes the union and
    says so in ``adjudication_note``. Unreviewed candidates are skipped and reported; unresolved
    disagreements are problems. The result is checked for split leakage.
    """
    candidates = list(candidates)
    resolved = {r.case_id: r for r in resolutions}
    reviews: dict[str, list[Review]] = {}
    for rs in review_sets:
        for r in rs:
            reviews.setdefault(r.case_id, []).append(r)

    cases: list[EvalCase] = []
    problems: list[str] = []
    unlabeled: list[str] = []
    unioned: list[str] = []
    for c in candidates:
        rs = reviews.get(c.case_id, [])
        final = resolved.get(c.case_id)
        if final is None:
            if not rs:
                unlabeled.append(c.case_id)
                continue
            if any(not _same_judgments(rs[0], r) for r in rs[1:]):
                problems.append(f"{c.case_id}: reviewers disagree and no resolution was provided")
                continue
            if any(not _same_reason(rs[0], r) for r in rs[1:]):
                problems.append(f"{c.case_id}: reviewers agree on the judgments but differ on the "
                                "primary reason, and no resolution was provided")
                continue
            final = rs[0]
        if final.reviewer == RESOLVER:
            labeled_by = ", ".join(r.reviewer for r in rs) or RESOLVER
            reviewed_by, note = RESOLVER, final.flag
            refs = list(final.evidence_refs)
        else:
            labeled_by = final.reviewer
            reviewed_by, note = (", ".join(r.reviewer for r in rs[1:]) or None), None
            refs = list(dict.fromkeys(ref for r in rs for ref in r.evidence_refs))
            if any(not _same_evidence(rs[0], r) for r in rs[1:]):
                cited = "; ".join(r.reviewer + ": " + ", ".join(r.evidence_refs) for r in rs)
                note = f"reviewers cited different evidence spans; the union was taken ({cited})"
                unioned.append(c.case_id)
        try:
            cases.append(EvalCase(input=c.input, gold=GoldLabel(
                case_id=c.case_id, relationship=derived_label(c, final),
                concept_same=final.concept_same, scope_same=final.scope_same,
                reasons=list(final.reasons), required_evidence_refs=refs,
                rationale=final.rationale, tags=list(c.tags), split=c.split, source=c.source,
                labeled_by=labeled_by, reviewed_by=reviewed_by, adjudication_note=note,
                guideline_version=GUIDELINE_VERSION,
            )))
        except ValueError as e:
            problems.append(f"{c.case_id}: {e}")
    problems += validate_cases(cases)
    return FinalizeResult(cases=cases, problems=problems, unlabeled=unlabeled,
                          evidence_unioned=unioned)
